import os
from PIL import Image
import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Colors
BG_DARK = RGBColor(17, 22, 29)        # #11161D
CARD_BG = RGBColor(26, 32, 42)        # #1A202A
CARD_BORDER = RGBColor(45, 55, 72)    # #2D3748
GOLD = RGBColor(255, 184, 0)          # #FFB800
EMERALD = RGBColor(0, 210, 106)       # #00D26A
CYAN = RGBColor(56, 189, 248)         # #38BDF8
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(148, 163, 184)        # #94A3B8
CARD_BG_ALT = RGBColor(22, 28, 38)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

IMG_DIR = r"g:\Projects\Project_1\docs\Img"
OUTPUT_PPTX = r"g:\Projects\Project_1\docs\MineMart_Pitch_Deck.pptx"

def create_slide_bg(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def add_header(slide, title_text, subtitle_text, gold_title=True):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.05))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = FONT_HEADING
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = GOLD if gold_title else WHITE
    
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = GRAY
    p2.space_before = Pt(3)

def add_image_fitted(slide, img_path, left, top, max_width, max_height, center=True):
    with Image.open(img_path) as im:
        orig_w, orig_h = im.size
    
    aspect = orig_w / orig_h
    box_aspect = max_width / max_height
    
    if aspect > box_aspect:
        fit_w = max_width
        fit_h = max_width / aspect
    else:
        fit_h = max_height
        fit_w = max_height * aspect
        
    actual_left = left + (max_width - fit_w) / 2 if center else left
    actual_top = top + (max_height - fit_h) / 2 if center else top
    
    pic = slide.shapes.add_picture(img_path, actual_left, actual_top, fit_w, fit_h)
    return pic

# ----------------- INITIALIZE PRESENTATION -----------------
prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# ==========================================================
# SLIDE 1: COVER & THÀNH VIÊN NHÓM
# ==========================================================
s1 = prs.slides.add_slide(blank_layout)
create_slide_bg(s1)

add_card(s1, Inches(0.8), Inches(0.8), Inches(6.8), Inches(5.9), bg_color=CARD_BG)

poly_path = os.path.join(IMG_DIR, "FPT-poly.png")
add_image_fitted(s1, poly_path, Inches(1.2), Inches(1.2), Inches(2.2), Inches(0.9), center=False)

tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.3), Inches(6.0), Inches(2.5))
tf1 = tb1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p = tf1.paragraphs[0]
p.text = "MINEMART"
p.font.name = FONT_HEADING
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = GOLD

p2 = tf1.add_paragraph()
p2.text = "POCKET WORKSHOP"
p2.font.name = FONT_HEADING
p2.font.size = Pt(26)
p2.font.bold = True
p2.font.color.rgb = EMERALD
p2.space_before = Pt(2)

p3 = tf1.add_paragraph()
p3.text = "Single-Screen Arcade Idle Tycoon (Voxel 3D)"
p3.font.name = FONT_BODY
p3.font.size = Pt(14)
p3.font.color.rgb = GRAY
p3.space_before = Pt(8)

add_card(s1, Inches(1.2), Inches(4.7), Inches(6.0), Inches(1.6), bg_color=CARD_BG_ALT, border_color=EMERALD)
tb_mem = s1.shapes.add_textbox(Inches(1.4), Inches(4.85), Inches(5.6), Inches(1.3))
tf_mem = tb_mem.text_frame
tf_mem.word_wrap = True
tf_mem.margin_left = tf_mem.margin_top = tf_mem.margin_right = tf_mem.margin_bottom = 0

pm1 = tf_mem.paragraphs[0]
pm1.text = "Tác giả: Lâm Tuyết Nhi (Miyuki)"
pm1.font.name = FONT_HEADING
pm1.font.size = Pt(16)
pm1.font.bold = True
pm1.font.color.rgb = WHITE

pm2 = tf_mem.add_paragraph()
pm2.text = "Vai trò: Trưởng nhóm / 3D Voxel Artist & Unity Developer"
pm2.font.name = FONT_BODY
pm2.font.size = Pt(13)
pm2.font.color.rgb = CYAN
pm2.space_before = Pt(4)

pm3 = tf_mem.add_paragraph()
pm3.text = "Đơn vị đào tạo: Cao đẳng FPT Polytechnic"
pm3.font.name = FONT_BODY
pm3.font.size = Pt(12)
pm3.font.color.rgb = GRAY
pm3.space_before = Pt(3)

add_card(s1, Inches(7.9), Inches(0.8), Inches(4.6), Inches(5.9), bg_color=CARD_BG)
poster_path = os.path.join(IMG_DIR, "MineMart_Poster_A2.png")
add_image_fitted(s1, poster_path, Inches(8.0), Inches(0.9), Inches(4.4), Inches(5.7))

# ==========================================================
# SLIDE 2: NGUỒN CẢM HƯNG & NĂNG LỰC BLOCKBENCH
# ==========================================================
s2 = prs.slides.add_slide(blank_layout)
create_slide_bg(s2)
add_header(s2, "NGUỒN CẢM HƯNG & THẾ MẠNH THỰC THI", 
           "Triết lý 'Thực học - Thực nghiệp' hội tụ cùng năng lực chế tác 3D Voxel Handcrafted")

add_card(s2, Inches(0.8), Inches(1.6), Inches(5.3), Inches(5.4), bg_color=CARD_BG)
add_image_fitted(s2, poly_path, Inches(1.1), Inches(1.8), Inches(2.6), Inches(1.0), center=False)

tb_s2_left = s2.shapes.add_textbox(Inches(1.1), Inches(2.9), Inches(4.7), Inches(3.8))
tf_s2_l = tb_s2_left.text_frame
tf_s2_l.word_wrap = True
tf_s2_l.margin_left = tf_s2_l.margin_top = tf_s2_l.margin_right = tf_s2_l.margin_bottom = 0

def add_bullet_point(tf, title, desc, title_color=GOLD):
    p_t = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
    p_t.text = "✦ " + title
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(14)
    p_t.font.bold = True
    p_t.font.color.rgb = title_color
    p_t.space_before = Pt(10)
    
    p_d = tf.add_paragraph()
    p_d.text = desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(12)
    p_d.font.color.rgb = GRAY
    p_d.space_before = Pt(2)

add_bullet_point(tf_s2_l, "Triết lý 'Thực học - Thực nghiệp'", 
                 "Không dừng lại ở một đồ án làm để lấy điểm số, em dốc 100% sức lực để kiểm tra giới hạn năng lực thực tế của bản thân.")
add_bullet_point(tf_s2_l, "Mục tiêu phát hành thực tế", 
                 "Kiên định đưa trò chơi lên Store để đón nhận phản hồi công tâm từ người chơi và thị trường xã hội bên ngoài.", title_color=EMERALD)
add_bullet_point(tf_s2_l, "Áp lực hóa động lực", 
                 "Biến tinh thần học tập năng động của sinh viên Poly thành ngọn lửa chế tác một sản phẩm hoàn thiện chỉn chu.", title_color=CYAN)

add_card(s2, Inches(6.4), Inches(1.6), Inches(6.1), Inches(5.4), bg_color=CARD_BG)

tb_s2_r_title = s2.shapes.add_textbox(Inches(6.7), Inches(1.8), Inches(5.5), Inches(0.6))
tf_s2_rt = tb_s2_r_title.text_frame
tf_s2_rt.word_wrap = True
tf_s2_rt.margin_left = tf_s2_rt.margin_top = tf_s2_rt.margin_right = tf_s2_rt.margin_bottom = 0
pr = tf_s2_rt.paragraphs[0]
pr.text = "100% 3D VOXEL HANDCRAFTED (BLOCKBENCH)"
pr.font.name = FONT_HEADING
pr.font.size = Pt(16)
pr.font.bold = True
pr.font.color.rgb = GOLD

blockbench_path = os.path.join(IMG_DIR, "blockbench_img.png")
add_image_fitted(s2, blockbench_path, Inches(6.7), Inches(2.4), Inches(5.5), Inches(3.0))

tb_s2_r_desc = s2.shapes.add_textbox(Inches(6.7), Inches(5.5), Inches(5.5), Inches(1.3))
tf_s2_rd = tb_s2_r_desc.text_frame
tf_s2_rd.word_wrap = True
tf_s2_rd.margin_left = tf_s2_rd.margin_top = tf_s2_rd.margin_right = tf_s2_rd.margin_bottom = 0
prd1 = tf_s2_rd.paragraphs[0]
prd1.text = "✔ Tự tay dựng mô hình & vẽ texture pixel từ nhân vật, quầy mỏ tới bàn phục vụ."
prd1.font.name = FONT_BODY
prd1.font.size = Pt(12)
prd1.font.color.rgb = WHITE
prd2 = tf_s2_rd.add_paragraph()
prd2.text = "✔ Tối ưu hóa Low-Poly chuẩn mobile, đảm bảo 60 FPS mượt mà không nóng máy."
prd2.font.name = FONT_BODY
prd2.font.size = Pt(12)
prd2.font.color.rgb = GRAY
prd2.space_before = Pt(3)

# ==========================================================
# SLIDE 3: BẢN SẮC GAMEPLAY — 1 MÀN HÌNH DỌC ĐỘNG
# ==========================================================
s3 = prs.slides.add_slide(blank_layout)
create_slide_bg(s3)
add_header(s3, "BẢN SẮC GAMEPLAY: HYBRID ARCADE IDLE", 
           "Trọn vẹn trên một màn hình dọc động — Cân bằng hoàn hảo giữa Can thiệp tay & Tự động hóa")

add_card(s3, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.4), bg_color=CARD_BG)
ingame_path = os.path.join(IMG_DIR, "ingame.png")
add_image_fitted(s3, ingame_path, Inches(0.95), Inches(1.75), Inches(3.3), Inches(5.1))

card_data_s3 = [
    {
        "title": "1. MÀN HÌNH DỌC ĐỘNG (SINGLE-SCREEN 2.5D)",
        "color": GOLD,
        "items": [
            "Bao quát toàn bộ xưởng sản xuất: Trạm mỏ → Quầy phục vụ khách/bàn → Tiền Tip.",
            "Thao tác bằng 1 ngón tay, góc nhìn Isometric trực quan, không chuyển cảnh gây nhiễu."
        ]
    },
    {
        "title": "2. CAN THIỆP TAY CÓ GIÁ TRỊ THẶNG DƯ (ACTIVE)",
        "color": EMERALD,
        "items": [
            "Kéo cần Joystick: Tốc độ đào x2 lần, nhặt tiền Tip trên đĩa nhận thêm +20% tiền thưởng.",
            "Cơ chế nhấc quầy: Cho phép người chơi tự do di dời và tái quy hoạch không gian xưởng."
        ]
    },
    {
        "title": "3. TỰ ĐỘNG HÓA THỎA MÃN (PURE-IDLE & AUTO-PILOT)",
        "color": CYAN,
        "items": [
            "Buông tay: Nhân vật tự hành làm việc cùng thợ mỏ với dáng đi lắc lư (waddling) cưng xỉu.",
            "Tự động tích lũy tài nguyên và lợi nhuận liên tục ngay cả khi người chơi rảnh tay."
        ]
    }
]

y_pos = 1.6
for c in card_data_s3:
    add_card(s3, Inches(4.7), Inches(y_pos), Inches(7.8), Inches(1.65), bg_color=CARD_BG)
    tb = s3.shapes.add_textbox(Inches(5.0), Inches(y_pos + 0.15), Inches(7.2), Inches(1.35))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = c["title"]
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = c["color"]
    
    for it in c["items"]:
        pi = tf.add_paragraph()
        pi.text = "• " + it
        pi.font.name = FONT_BODY
        pi.font.size = Pt(11.5)
        pi.font.color.rgb = WHITE if "x2" in it or "Single" in it else GRAY
        pi.space_before = Pt(3)
        
    y_pos += 1.85

# ==========================================================
# SLIDE 4: ỨNG DỤNG AI: TỰ ĐỘNG HÓA CÔNG CỤ PHÁT TRIỂN (TOOLING)
# ==========================================================
s4 = prs.slides.add_slide(blank_layout)
create_slide_bg(s4)
add_header(s4, "ỨNG DỤNG AI: TỰ ĐỘNG HÓA CÔNG CỤ PHÁT TRIỂN (TOOLING)", 
           "Xây dựng bộ công cụ Unity Editor chuyên biệt — Tự động hóa quy trình, tăng tốc độ sản xuất game")

# 3 Columns matching exactly the 3 images
card_data_s4 = [
    {
        "x": 0.8,
        "w": 3.7,
        "title": "1. TOP BAR MENU SYSTEM",
        "subtitle": "Thanh công cụ MineMart > UI",
        "color": GOLD,
        "img": "AI1 (2).png",
        "bullets": [
            "Tích hợp sẵn vào Top Menu của Unity Editor.",
            "1-Click sinh toàn bộ Settings Dialog & Startup Flow.",
            "Tự động hóa nút Reset Progress & đổi màu Loading.",
            "Giảm 80% thời gian cấu hình UI thủ công."
        ]
    },
    {
        "x": 4.8,
        "w": 3.7,
        "title": "2. HIERARCHY CONTEXT MENU",
        "subtitle": "Menu chuột phải MineMart UI",
        "color": EMERALD,
        "img": "AI1 (3).png",
        "bullets": [
            "Menu chuột phải thông minh ngay trong Hierarchy.",
            "Tự sinh Prefab thẻ mỏ & Plot Stall Shop Bubble.",
            "Tạo cây nâng cấp Ngọc (EUGT) & Map Upgrade Panel.",
            "Chuẩn hóa cấu trúc Anchor & Canvas chuẩn xác."
        ]
    },
    {
        "x": 8.8,
        "w": 3.7,
        "title": "3. CUSTOM INSPECTOR TOOLING",
        "subtitle": "Công cụ Milestone Slider trực quan",
        "color": CYAN,
        "img": "AI1 (1).png",
        "bullets": [
            "Nút bấm thông minh sinh Slider ngay trong Scene View.",
            "Tự động liên kết Slider, Fill Image & Container.",
            "Cho phép tùy chỉnh màu sắc & vị trí thời gian thực.",
            "Phát hiện và cảnh báo thiếu liên kết tự động."
        ]
    }
]

for col in card_data_s4:
    add_card(s4, Inches(col["x"]), Inches(1.55), Inches(col["w"]), Inches(4.65), bg_color=CARD_BG)
    
    # Title & Subtitle box
    tb = s4.shapes.add_textbox(Inches(col["x"] + 0.2), Inches(1.68), Inches(col["w"] - 0.4), Inches(0.55))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = col["title"]
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = col["color"]
    
    psub = tf.add_paragraph()
    psub.text = col["subtitle"]
    psub.font.name = FONT_BODY
    psub.font.size = Pt(10)
    psub.font.color.rgb = GRAY
    psub.space_before = Pt(1)
    
    # Image in middle
    ai_img_path = os.path.join(IMG_DIR, col["img"])
    add_image_fitted(s4, ai_img_path, Inches(col["x"] + 0.2), Inches(2.3), Inches(col["w"] - 0.4), Inches(1.65))
    
    # Bullets below
    tb_b = s4.shapes.add_textbox(Inches(col["x"] + 0.2), Inches(4.05), Inches(col["w"] - 0.4), Inches(2.05))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_top = tf_b.margin_right = tf_b.margin_bottom = 0
    
    first = True
    for b in col["bullets"]:
        pb = tf_b.add_paragraph() if not first else tf_b.paragraphs[0]
        first = False
        pb.text = "• " + b
        pb.font.name = FONT_BODY
        pb.font.size = Pt(10.5)
        pb.font.color.rgb = WHITE
        pb.space_before = Pt(3)

# Bottom Summary Banner for Slide 4
add_card(s4, Inches(0.8), Inches(6.32), Inches(11.7), Inches(0.68), bg_color=CARD_BG_ALT, border_color=EMERALD)
tb_bann = s4.shapes.add_textbox(Inches(1.0), Inches(6.38), Inches(11.3), Inches(0.55))
tf_bann = tb_bann.text_frame
tf_bann.word_wrap = True
tf_bann.margin_left = tf_bann.margin_top = tf_bann.margin_right = tf_bann.margin_bottom = 0
pbann = tf_bann.paragraphs[0]
pbann.text = "💡 Dấu ấn AI: Tự động hóa bộ công cụ Editor • Tối ưu thuật toán điều phối FIFO O(1) • Tính toán chuyển động nhún nhảy Procedural Waddle"
pbann.font.name = FONT_BODY
pbann.font.size = Pt(11.5)
pbann.font.bold = True
pbann.font.color.rgb = GOLD

# ==========================================================
# SLIDE 5: LỘ TRÌNH PHÁT HÀNH & MÔ HÌNH KIẾM TIỀN
# ==========================================================
s5 = prs.slides.add_slide(blank_layout)
create_slide_bg(s5)
add_header(s5, "LỘ TRÌNH PHÁT HÀNH & TIỀM NĂNG THƯƠNG MẠI", 
           "Kế hoạch hành động cụ thể để đưa MineMart từ đồ án thành tựa game thương mại hoàn chỉnh")

pillars = [
    {
        "title": "1. KIẾN TRÚC DATA-DRIVEN & EXPANSION",
        "color": GOLD,
        "desc": "Tách toàn bộ chỉ số ra ScriptableObjects, dễ dàng mở rộng thêm các Biomes mới: Sa Mạc, Đảo Băng, Rừng Nhiệt Đới."
    },
    {
        "title": "2. HỆ THỐNG TRANG BỊ & THÚ CƯNG",
        "color": EMERALD,
        "desc": "Bổ sung rương mở khóa cuốc mỏ gia tăng tốc độ đào và pet đồng hành hỗ trợ tăng tốc độ vận chuyển hàng."
    },
    {
        "title": "3. MÔ HÌNH KIẾM TIỀN HYBRID CASUAL",
        "color": CYAN,
        "desc": "• Rewarded Ads: Xem video x2 doanh thu 5 phút / gọi Siêu Thợ Mỏ.\n• In-App Purchase: Gói No-Ads vĩnh viễn và Skin độc quyền (Chibi Voxel / Áo FPT Poly)."
    }
]

y_pos_s5 = 1.6
for pil in pillars:
    add_card(s5, Inches(0.8), Inches(y_pos_s5), Inches(7.5), Inches(1.65), bg_color=CARD_BG)
    tb = s5.shapes.add_textbox(Inches(1.1), Inches(y_pos_s5 + 0.15), Inches(6.9), Inches(1.35))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p = tf.paragraphs[0]
    p.text = pil["title"]
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = pil["color"]
    
    pd = tf.add_paragraph()
    pd.text = pil["desc"]
    pd.font.name = FONT_BODY
    pd.font.size = Pt(11.5)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(3)
    
    y_pos_s5 += 1.85

add_card(s5, Inches(8.6), Inches(1.6), Inches(3.9), Inches(5.4), bg_color=CARD_BG)

icon_path = os.path.join(IMG_DIR, "MineMart_Icon.jpg")
add_image_fitted(s5, icon_path, Inches(9.1), Inches(1.9), Inches(2.9), Inches(2.9))

tb_cta = s5.shapes.add_textbox(Inches(8.8), Inches(4.9), Inches(3.5), Inches(1.9))
tf_cta = tb_cta.text_frame
tf_cta.word_wrap = True
tf_cta.margin_left = tf_cta.margin_top = tf_cta.margin_right = tf_cta.margin_bottom = 0

pcta1 = tf_cta.paragraphs[0]
pcta1.text = "MỤC TIÊU PHÁT HÀNH"
pcta1.alignment = PP_ALIGN.CENTER
pcta1.font.name = FONT_HEADING
pcta1.font.size = Pt(15)
pcta1.font.bold = True
pcta1.font.color.rgb = GOLD

pcta2 = tf_cta.add_paragraph()
pcta2.text = "Google Play Store & itch.io"
pcta2.alignment = PP_ALIGN.CENTER
pcta2.font.name = FONT_BODY
pcta2.font.size = Pt(13)
pcta2.font.bold = True
pcta2.font.color.rgb = WHITE
pcta2.space_before = Pt(3)

pcta3 = tf_cta.add_paragraph()
pcta3.text = "Cảm ơn quý Thầy Cô đã lắng nghe!"
pcta3.alignment = PP_ALIGN.CENTER
pcta3.font.name = FONT_BODY
pcta3.font.size = Pt(12)
pcta3.font.color.rgb = EMERALD
pcta3.space_before = Pt(8)

# Save
prs.save(OUTPUT_PPTX)
print("SUCCESS: PowerPoint saved to", OUTPUT_PPTX)
